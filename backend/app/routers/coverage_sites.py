"""
Coverage sites and contacts outside the net.

For coverage drills such as an ARES Simulated Emergency Test, where the
question is which places can reach each other, not which operators:

  /nets/{id}/coverage-sites     the sites the drill is testing (net manager)
  /nets/{id}/coverage-contacts  stations reached that never checked in
  /nets/{id}/coverage-access    which repeaters and nodes a station got into
  /nets/{id}/site-coverage      the read-time rollup (services/site_coverage.py)

Reads are open to guests, like the check-in list and the "can hear" reports,
with one exception: an imported list of possible shelters or distribution
points is not public. Anyone but net staff sees only the sites in use
(site_coverage.coverage_site_in_use), and never a site's notes or details.
Every write broadcasts can_hear_changed, the one event that already makes
every viewer refetch coverage.

  POST /nets/{id}/coverage-sites/read-file  reads an EMA's .xlsx or .csv
  site list for the import, and saves nothing (net editors)
"""
import csv
import datetime as dt
import io
import json
from typing import List, Optional, Set

from fastapi import APIRouter, Depends, File, Form, HTTPException, UploadFile
from sqlalchemy import delete, select, update
from sqlalchemy.ext.asyncio import AsyncSession
from sqlalchemy.orm import selectinload

from app.database import get_db
from app.dependencies import get_current_user, get_current_user_optional
from app.models import CheckIn, CoverageAccess, CoverageContact, CoverageSite, DrillEvent, Net, User, net_frequencies
from app.permissions import may_edit_net, may_run_drill, net_access
from app.routers.can_hear import may_report_coverage_for, validate_coverage_frequency
from app.schemas import (
    CoverageAccessResponse, CoverageAccessSave, CoverageContactResponse, CoverageContactsSave, CoverageSiteItem,
    CoverageSiteResponse, SiteCoverageResponse, SiteFileResponse,
)
from app.services.site_coverage import coverage_site_in_use, load_site_coverage

router = APIRouter(prefix="/nets", tags=["coverage-sites"])

MAX_SITES = 500
# The import reads at most this much of a spreadsheet: an EMA's list is a few
# hundred rows, and the caps keep a malformed or hostile file from costing
# more than that.
MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_FILE_ROWS = 2000
MAX_FILE_COLUMNS = 40


async def _load_net(db: AsyncSession, net_id: int) -> Net:
    net = (await db.execute(
        select(Net).options(selectinload(Net.frequencies)).where(Net.id == net_id)
    )).scalar_one_or_none()
    if not net:
        raise HTTPException(status_code=404, detail="Net not found")
    return net


async def _broadcast_coverage_changed(net_id: int, **data) -> None:
    from app.main import manager
    await manager.broadcast({
        "type": "can_hear_changed",
        "data": {"net_id": net_id, **data},
        "timestamp": dt.datetime.utcnow().isoformat(),
    }, net_id)


async def _sites(db: AsyncSession, net_id: int) -> List[CoverageSite]:
    return (await db.execute(
        select(CoverageSite).where(CoverageSite.net_id == net_id).order_by(CoverageSite.sort_order, CoverageSite.id)
    )).scalars().all()


async def may_see_every_site(db: AsyncSession, net: Net, user: Optional[User]) -> bool:
    """Net staff see a net's whole site list and every site's notes and
    details; everyone else sees only the sites in use. Staff is whoever may
    run the drill or edit the net."""
    if user is None:
        return False
    access = await net_access(db, net, user)
    return may_run_drill(access) or may_edit_net(access)


async def checked_in_site_ids(db: AsyncSession, net_id: int) -> Set[int]:
    return set((await db.execute(
        select(CheckIn.coverage_site_id).where(CheckIn.net_id == net_id, CheckIn.coverage_site_id.is_not(None)).distinct()
    )).scalars().all())


def _site_response(site: CoverageSite, in_use: bool, staff: bool) -> CoverageSiteResponse:
    return CoverageSiteResponse(
        id=site.id, net_id=site.net_id, name=site.name, location=site.location, site_type=site.site_type,
        is_primary=bool(site.is_primary), sort_order=site.sort_order, frequency_id=site.frequency_id,
        address=site.address, latitude=site.latitude, longitude=site.longitude,
        is_supported=bool(site.is_supported), in_use=in_use,
        notes=site.notes if staff else None,
        details=(json.loads(site.details) if site.details else None) if staff else None,
    )


async def _site_list(db: AsyncSession, net: Net, user: Optional[User]) -> List[CoverageSiteResponse]:
    staff = await may_see_every_site(db, net, user)
    used = await checked_in_site_ids(db, net.id)
    out = []
    for site in await _sites(db, net.id):
        in_use = coverage_site_in_use(site, used)
        if in_use or staff:
            out.append(_site_response(site, in_use, staff))
    return out


@router.get("/{net_id}/coverage-sites", response_model=List[CoverageSiteResponse])
async def list_coverage_sites(
    net_id: int,
    current_user: Optional[User] = Depends(get_current_user_optional),
    db: AsyncSession = Depends(get_db),
):
    """The net's sites, in the manager's order: every site for net staff,
    only the sites in use for everyone else."""
    net = await _load_net(db, net_id)
    return await _site_list(db, net, current_user)


@router.put("/{net_id}/coverage-sites", response_model=List[CoverageSiteResponse])
async def save_coverage_sites(
    net_id: int,
    items: List[CoverageSiteItem],
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Replace the net's site list. Items with an id keep that site, so the
    check-ins already placed at it stay there; a site left out is removed and
    its check-ins are no longer at any site. Same rule as editing the net."""
    net = await _load_net(db, net_id)
    if not may_edit_net(await net_access(db, net, current_user)):
        raise HTTPException(status_code=403, detail="Not authorized to change this net's sites")
    if len(items) > MAX_SITES:
        raise HTTPException(status_code=400, detail=f"A net can have at most {MAX_SITES} sites")
    # A name can repeat in another town (every town has a "Fire Station"),
    # so a site is told apart by its name and place together.
    names = [(i.name.strip().lower(), (i.location or '').strip().lower()) for i in items]
    if len(set(names)) != len(names):
        raise HTTPException(status_code=400, detail="Two sites have the same name and place")

    existing = {s.id: s for s in await _sites(db, net_id)}
    unknown = {i.id for i in items if i.id is not None} - set(existing)
    if unknown:
        raise HTTPException(status_code=400, detail="One or more sites don't belong to this net")

    net_frequency_ids = {f.id for f in net.frequencies}
    if any(i.frequency_id is not None and i.frequency_id not in net_frequency_ids for i in items):
        raise HTTPException(status_code=400, detail="Frequency does not belong to this net")

    kept_ids = {i.id for i in items if i.id is not None}
    removed_ids = set(existing) - kept_ids
    if removed_ids:
        # SQLite doesn't enforce the ON DELETE SET NULL, so do it here.
        await db.execute(
            update(CheckIn).where(CheckIn.coverage_site_id.in_(removed_ids)).values(coverage_site_id=None)
        )
        await db.execute(delete(CoverageAccess).where(CoverageAccess.target_site_id.in_(removed_ids)))
        await db.execute(update(DrillEvent).where(DrillEvent.site_id.in_(removed_ids)).values(site_id=None))
        for site_id in removed_ids:
            await db.delete(existing[site_id])

    for order, item in enumerate(items):
        site = existing.get(item.id) if item.id is not None else None
        if site is None:
            site = CoverageSite(net_id=net_id)
            db.add(site)
        site.name = item.name.strip()
        site.location = (item.location or '').strip() or None
        site.site_type = item.site_type
        # A repeater or node is something sites reach, never a place others
        # are graded against, and only it is tied to a frequency.
        infrastructure = item.site_type == 'infrastructure'
        site.is_primary = item.is_primary and not infrastructure
        site.frequency_id = item.frequency_id if infrastructure else None
        site.address = (item.address or '').strip() or None
        site.latitude = item.latitude if item.latitude is not None and item.longitude is not None else None
        site.longitude = item.longitude if site.latitude is not None else None
        site.notes = (item.notes or '').strip() or None
        site.details = json.dumps(item.details) if item.details else None
        site.is_supported = item.is_supported
        site.sort_order = order

    await db.flush()
    infrastructure_ids = [s.id for s in await _sites(db, net_id) if s.site_type == 'infrastructure']
    if infrastructure_ids:
        await db.execute(
            update(CheckIn).where(CheckIn.coverage_site_id.in_(infrastructure_ids)).values(coverage_site_id=None)
        )
    await db.commit()
    await _broadcast_coverage_changed(net_id, sites_changed=True)
    return await _site_list(db, net, current_user)


def _cell_text(value) -> str:
    """A spreadsheet cell as text. Whole-number floats lose their ".0" (a
    capacity of 150 reads as 150), and nothing is ever evaluated."""
    if value is None:
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (dt.datetime, dt.date)):
        return value.isoformat()
    return str(value).strip()


def _read_xlsx(data: bytes, sheet: Optional[str]):
    # defusedxml must be importable: openpyxl uses it when present to refuse
    # the XML entity attacks (billion laughs, quadratic blowup) that it
    # doesn't guard against on its own. Fail closed rather than parse
    # without it.
    import defusedxml  # noqa: F401
    from openpyxl import load_workbook
    # data_only reads the value Excel last saved, never the formula.
    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        names = workbook.sheetnames
        chosen = sheet if sheet in names else names[0]
        rows, truncated = [], False
        for row in workbook[chosen].iter_rows(values_only=True):
            if len(rows) >= MAX_FILE_ROWS:
                truncated = True
                break
            rows.append([_cell_text(v) for v in row[:MAX_FILE_COLUMNS]])
        return names, chosen, rows, truncated
    finally:
        workbook.close()


def _read_csv(data: bytes):
    try:
        text = data.decode('utf-8-sig')
    except UnicodeDecodeError:
        text = data.decode('latin-1')
    rows, truncated = [], False
    for row in csv.reader(io.StringIO(text)):
        if len(rows) >= MAX_FILE_ROWS:
            truncated = True
            break
        rows.append([c.strip() for c in row[:MAX_FILE_COLUMNS]])
    return [], None, rows, truncated


@router.post("/{net_id}/coverage-sites/read-file", response_model=SiteFileResponse)
async def read_site_file(
    net_id: int,
    file: UploadFile = File(...),
    sheet: Optional[str] = Form(None),
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Read an .xlsx or .csv list of sites for the import dialog. Returns
    its cells as text, with blank rows dropped; matching columns, choosing
    the type, and saving all happen afterward, through the site list's own
    save. The file is never stored. Same rule as editing the net's sites."""
    net = await _load_net(db, net_id)
    if not may_edit_net(await net_access(db, net, current_user)):
        raise HTTPException(status_code=403, detail="Not authorized to change this net's sites")
    data = await file.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise HTTPException(status_code=413, detail="That file is over 5 MB")
    name = (file.filename or '').lower()
    try:
        if name.endswith('.xlsx') or data[:2] == b'PK':
            sheets, chosen, rows, truncated = _read_xlsx(data, sheet)
        elif name.endswith('.csv') or name.endswith('.txt'):
            sheets, chosen, rows, truncated = _read_csv(data)
        else:
            raise HTTPException(status_code=400, detail="Choose an Excel workbook (.xlsx) or a CSV file")
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(status_code=400, detail="That file couldn't be read as a spreadsheet")
    rows = [r for r in rows if any(c for c in r)]
    return SiteFileResponse(sheets=sheets, sheet=chosen, rows=rows, truncated=truncated)


@router.get("/{net_id}/coverage-contacts", response_model=List[CoverageContactResponse])
async def list_coverage_contacts(net_id: int, db: AsyncSession = Depends(get_db)):
    """Every station reached outside the net, oldest first."""
    await _load_net(db, net_id)
    rows = (await db.execute(
        select(CoverageContact).where(CoverageContact.net_id == net_id).order_by(CoverageContact.id)
    )).scalars().all()
    return [CoverageContactResponse.model_validate(c) for c in rows]


@router.put("/{net_id}/coverage-contacts", response_model=List[CoverageContactResponse])
async def save_coverage_contacts(
    net_id: int,
    payload: CoverageContactsSave,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Replace every outside contact one check-in has recorded, the same
    whole-set save the "can hear" dialog uses. Same who-may rule too."""
    net = await _load_net(db, net_id)
    if not net.propagation_logging_enabled:
        raise HTTPException(status_code=403, detail="Station-to-station coverage logging is not enabled for this net")

    reporter = (await db.execute(
        select(CheckIn).where(CheckIn.id == payload.reporter_check_in_id)
    )).scalar_one_or_none()
    if not reporter:
        raise HTTPException(status_code=404, detail="Check-in not found")
    if reporter.net_id != net_id:
        raise HTTPException(status_code=400, detail="Check-in does not belong to this net")
    if not await may_report_coverage_for(db, net, reporter, current_user):
        raise HTTPException(status_code=403, detail="Not authorized to record coverage reports")
    for item in payload.contacts:
        await validate_coverage_frequency(db, net, item.frequency_id)

    existing = (await db.execute(
        select(CoverageContact).where(CoverageContact.reporter_check_in_id == reporter.id)
    )).scalars().all()
    for row in existing:
        await db.delete(row)
    now = dt.datetime.utcnow()
    for item in payload.contacts:
        db.add(CoverageContact(
            net_id=net_id,
            reporter_check_in_id=reporter.id,
            frequency_id=item.frequency_id,
            callsign=item.callsign,
            station=(item.station or '').strip() or None,
            county=(item.county or '').strip() or None,
            state=(item.state or '').strip() or None,
            two_way=item.two_way,
            reported_by_user_id=current_user.id,
            reported_at=now,
        ))
    await db.commit()
    await _broadcast_coverage_changed(net_id, reporter_check_in_id=reporter.id)

    rows = (await db.execute(
        select(CoverageContact).where(CoverageContact.reporter_check_in_id == reporter.id).order_by(CoverageContact.id)
    )).scalars().all()
    return [CoverageContactResponse.model_validate(c) for c in rows]


@router.get("/{net_id}/coverage-access", response_model=List[CoverageAccessResponse])
async def list_coverage_access(net_id: int, db: AsyncSession = Depends(get_db)):
    """Every repeater/node result on the net."""
    await _load_net(db, net_id)
    rows = (await db.execute(
        select(CoverageAccess).where(CoverageAccess.net_id == net_id).order_by(CoverageAccess.id)
    )).scalars().all()
    return [CoverageAccessResponse.model_validate(a) for a in rows]


@router.put("/{net_id}/coverage-access", response_model=List[CoverageAccessResponse])
async def save_coverage_access(
    net_id: int,
    payload: CoverageAccessSave,
    current_user: User = Depends(get_current_user),
    db: AsyncSession = Depends(get_db),
):
    """Replace every repeater/node result one check-in has: reached, couldn't
    reach, and an optional note each. A target left out wasn't tried. Same
    who-may rule as "can hear" reports."""
    net = await _load_net(db, net_id)
    if not net.propagation_logging_enabled:
        raise HTTPException(status_code=403, detail="Station-to-station coverage logging is not enabled for this net")
    reporter = (await db.execute(
        select(CheckIn).where(CheckIn.id == payload.reporter_check_in_id)
    )).scalar_one_or_none()
    if not reporter:
        raise HTTPException(status_code=404, detail="Check-in not found")
    if reporter.net_id != net_id:
        raise HTTPException(status_code=400, detail="Check-in does not belong to this net")
    if not await may_report_coverage_for(db, net, reporter, current_user):
        raise HTTPException(status_code=403, detail="Not authorized to record coverage reports")

    target_ids = [r.target_site_id for r in payload.results]
    if len(set(target_ids)) != len(target_ids):
        raise HTTPException(status_code=400, detail="Each repeater or node can only have one result")
    if target_ids:
        valid = set((await db.execute(
            select(CoverageSite.id).where(
                CoverageSite.id.in_(target_ids), CoverageSite.net_id == net_id,
                CoverageSite.site_type == 'infrastructure',
            )
        )).scalars().all())
        if valid != set(target_ids):
            raise HTTPException(status_code=400, detail="One or more of those aren't this net's repeaters or nodes")

    await db.execute(delete(CoverageAccess).where(CoverageAccess.reporter_check_in_id == reporter.id))
    now = dt.datetime.utcnow()
    for r in payload.results:
        db.add(CoverageAccess(
            net_id=net_id, reporter_check_in_id=reporter.id, target_site_id=r.target_site_id,
            reached=r.reached, note=(r.note or '').strip() or None,
            reported_by_user_id=current_user.id, reported_at=now,
        ))
    await db.commit()
    await _broadcast_coverage_changed(net_id, reporter_check_in_id=reporter.id)
    rows = (await db.execute(
        select(CoverageAccess).where(CoverageAccess.reporter_check_in_id == reporter.id).order_by(CoverageAccess.id)
    )).scalars().all()
    return [CoverageAccessResponse.model_validate(a) for a in rows]


@router.get("/{net_id}/site-coverage", response_model=SiteCoverageResponse)
async def get_site_coverage(net_id: int, db: AsyncSession = Depends(get_db)):
    """Coverage by site: which sites reached which, on what, and whether each
    reaches the primary sites directly, through a relay, or not at all."""
    net = await _load_net(db, net_id)
    return await load_site_coverage(db, net)
